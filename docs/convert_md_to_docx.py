import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def md_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()

    # ตั้งค่าขอบหน้ากระดาษ
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'TH Sarabun New'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # กรองตัวคั่น |---|---|
        filtered = []
        for r in table_rows:
            if re.match(r'^\s*\|?(\s*:?-+:?\s*\|)+\s*$', r):
                continue
            cols = [c.strip() for c in r.strip().strip('|').split('|')]
            if cols and any(c for c in cols):
                filtered.append(cols)

        if not filtered:
            table_rows = []
            in_table = False
            return

        max_cols = max(len(r) for r in filtered)
        table = doc.add_table(rows=len(filtered), cols=max_cols)
        table.autofit = True

        for i, row in enumerate(filtered):
            is_header = (i == 0)
            for j in range(max_cols):
                cell = table.cell(i, j)
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                text = row[j] if j < len(row) else ""
                
                # Format cell text
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                # clean Markdown link
                clean_text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
                clean_text = clean_text.replace('<br>', '\n')
                
                run = p.add_run(clean_text)
                run.font.name = 'TH Sarabun New'
                run.font.size = Pt(13)
                
                if is_header:
                    set_cell_background(cell, "0B2046") # สีกรมท่า
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        doc.add_paragraph() # space after table
        table_rows = []
        in_table = False

    for line in lines:
        sline = line.rstrip()

        # Code block
        if sline.startswith("```"):
            if in_code_block:
                # end code block
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                code_text = "\n".join(code_lines)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                code_lines = []
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(sline)
            continue

        # Table detection
        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_rows.append(sline)
            continue
        elif in_table:
            flush_table()

        # Headings
        if sline.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(sline[2:])
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0B, 0x20, 0x46)
        elif sline.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sline[3:])
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif sline.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[4:])
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif sline.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline[5:])
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif sline.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[2:])
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(13)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif sline.strip().startswith("- ") or sline.strip().startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            clean_item = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', sline.strip()[2:])
            run = p.add_run(clean_item)
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(14)
        elif re.match(r'^\d+\.\s', sline.strip()):
            m = re.match(r'^(\d+\.\s)(.*)', sline.strip())
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            clean_item = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', m.group(2))
            run = p.add_run(clean_item)
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(14)
        elif sline.strip() == "---":
            # horizontal line / separator
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        elif sline.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            clean_text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', sline)
            run = p.add_run(clean_text)
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(14)

    if in_table:
        flush_table()

    doc.save(docx_path)
    print("SUCCESS: File saved to", docx_path)

if __name__ == "__main__":
    md_to_docx(r"C:\Project\HRMS\docs\project-guide-full.md", r"C:\Project\HRMS\docs\project-guide-full.docx")
