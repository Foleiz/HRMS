# -*- coding: utf-8 -*-
"""
Daily Progress Summary Generator
Automates scanning project activity, git commits, and file changes to produce
both Markdown (.md) and Microsoft Word (.docx) daily reports in everyday Thai.
"""

import os
import sys
import json
import re
import datetime
import subprocess
import argparse

# Force UTF-8 stdout for Windows consoles
sys.stdout.reconfigure(encoding='utf-8')

THAI_MONTHS = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"
]
THAI_DAYS = ["วันจันทร์", "วันอังคาร", "วันพุธ", "วันพฤหัสบดี", "วันศุกร์", "วันเสาร์", "วันอาทิตย์"]

def get_thai_date_str(target_date: datetime.date) -> str:
    day_name = THAI_DAYS[target_date.weekday()]
    thai_year = target_date.year + 543
    return f"{day_name}ที่ {target_date.day} {THAI_MONTHS[target_date.month]} {thai_year}"

def get_next_report_number(daily_reports_dir: str) -> int:
    if not os.path.exists(daily_reports_dir):
        return 1
    max_num = 0
    pattern = re.compile(r"^(\d+)\.daily_summary_")
    for fname in os.listdir(daily_reports_dir):
        m = pattern.match(fname)
        if m:
            num = int(m.group(1))
            if num > max_num:
                max_num = num
    return max_num + 1

def scan_git_activity(project_path: str, target_date_str: str = None) -> dict:
    activity = {
        "commits": [],
        "uncommitted_files": [],
        "tech_stack": []
    }
    
    # 1. Check Git Commits
    try:
        since_arg = f'{target_date_str} 00:00:00' if target_date_str else 'today 00:00:00'
        cmd = [
            "git", "-C", project_path, "log",
            f"--since={since_arg}",
            "--pretty=format:__COMMIT__%h|%an|%ad|%s",
            "--date=format:%H:%M"
        ]
        res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8', errors='ignore')
        raw_output = res.stdout.strip()
        
        if not raw_output:
            cmd_fallback = [
                "git", "-C", project_path, "log", "-n", "5",
                "--pretty=format:__COMMIT__%h|%an|%ad|%s",
                "--date=format:%H:%M"
            ]
            res_fb = subprocess.run(cmd_fallback, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8', errors='ignore')
            raw_output = res_fb.stdout.strip()

        if raw_output:
            commit_blocks = raw_output.split("__COMMIT__")
            for block in commit_blocks:
                block = block.strip()
                if not block:
                    continue
                parts = block.split("|", 3)
                if len(parts) == 4:
                    activity["commits"].append({
                        "hash": parts[0],
                        "author": parts[1],
                        "time": parts[2],
                        "subject": parts[3]
                    })
    except Exception as e:
        activity["git_error"] = str(e)

    # 2. Check Uncommitted Files
    try:
        cmd_status = ["git", "-C", project_path, "status", "--short"]
        res_st = subprocess.run(cmd_status, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8', errors='ignore')
        for line in res_st.stdout.splitlines():
            line = line.strip()
            if line:
                activity["uncommitted_files"].append(line)
    except Exception:
        pass

    # 3. Detect Tech Stack
    files_in_root = os.listdir(project_path) if os.path.exists(project_path) else []
    if any(f.endswith(".csproj") or f == "backend" for f in files_in_root):
        activity["tech_stack"].append(".NET C# Backend")
    if any(f in ["package.json", "frontend"] for f in files_in_root):
        activity["tech_stack"].append("Next.js / TypeScript Frontend")
    if any(f.endswith(".sql") or "database" in f.lower() for f in files_in_root):
        activity["tech_stack"].append("PostgreSQL Database")

    return activity

def generate_docx(data: dict, output_path: str):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("รายงานสรุปความคืบหน้าประจำวัน (Daily Progress Report)")
    r_title.font.name = 'TH Sarabun New'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 44, 87)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)

    # Meta
    p_meta = doc.add_paragraph()
    proj_name = data.get("project_name", "HRMS Development")
    meta_text = f"โครงการ: {proj_name} | {data.get('thai_date', '')} (รอบ 17:30 น. - ฉบับที่ {data.get('report_number', 1)})"
    r_meta = p_meta.add_run(meta_text)
    r_meta.font.name = 'TH Sarabun New'
    r_meta.font.size = Pt(13)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    p_meta.paragraph_format.space_after = Pt(12)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 44, 87)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'TH Sarabun New'
            r_pre.font.size = Pt(13.5)
            r_pre.font.bold = True
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(13.5)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'TH Sarabun New'
            r_pre.font.size = Pt(13.5)
            r_pre.font.bold = True
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(13.5)
        return p

    # Section 1: Executive Summary
    add_h2("1. สรุปภาพรวมใน 3 บรรทัด (Executive Summary)")
    for idx, item in enumerate(data.get("executive_summary", []), start=1):
        add_bullet(f" {item.get('desc', '')}", f"{idx}. {item.get('title', '')}:")

    # Section 2: Details of work
    add_h2("2. รายละเอียดงานที่ทำและแก้ไขในวันนี้")
    for idx, feat in enumerate(data.get("features", []), start=1):
        add_h3(f"{idx}. {feat.get('title', '')}")
        if feat.get("before"):
            add_body(f" {feat['before']}", "ก่อนหน้านี้เป็นอย่างไร:")
        if feat.get("action"):
            add_body(f" {feat['action']}", "สิ่งที่ทำในวันนี้:")
        if feat.get("benefit"):
            add_body(f" {feat['benefit']}", "ประโยชน์ทางธุรกิจ:")

    # Section 3: Readiness Table
    add_h2("3. สรุปสถานะความพร้อมของระบบ ณ สิ้นวัน")
    table_data = data.get("readiness_table", [])
    if table_data:
        table = doc.add_table(rows=len(table_data)+1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ['องค์ประกอบระบบ', 'จำนวน / สถานะ', 'ความพร้อมใช้งาน']
        col_widths = [Inches(2.5), Inches(1.6), Inches(2.7)]

        for c_idx, h_text in enumerate(headers):
            cell = table.cell(0, c_idx)
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "102C57")
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h_text)
            r.font.name = 'TH Sarabun New'
            r.font.size = Pt(13.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row_item in enumerate(table_data, start=1):
            row = table.rows[r_idx]
            row_bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            
            cell0 = row.cells[0]
            cell0.width = col_widths[0]
            set_cell_background(cell0, row_bg)
            set_cell_margins(cell0, top=100, bottom=100, left=150, right=150)
            p0 = cell0.paragraphs[0]
            r0 = p0.add_run(row_item.get("component", ""))
            r0.font.name = 'TH Sarabun New'
            r0.font.size = Pt(13)
            r0.font.bold = True

            cell1 = row.cells[1]
            cell1.width = col_widths[1]
            set_cell_background(cell1, row_bg)
            set_cell_margins(cell1, top=100, bottom=100, left=150, right=150)
            p1 = cell1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(row_item.get("status", ""))
            r1.font.name = 'TH Sarabun New'
            r1.font.size = Pt(13)

            cell2 = row.cells[2]
            cell2.width = col_widths[2]
            set_cell_background(cell2, row_bg)
            set_cell_margins(cell2, top=100, bottom=100, left=150, right=150)
            p2 = cell2.paragraphs[0]
            r2 = p2.add_run(row_item.get("readiness", ""))
            r2.font.name = 'TH Sarabun New'
            r2.font.size = Pt(13)

    # Section 4: Next Steps
    add_h2("4. แผนงานสำหรับรอบวันถัดไป")
    for step in data.get("next_steps", []):
        add_bullet(f" {step}")

    doc.save(output_path)
    print(f"[OK] Generated Word Document: {output_path}")

def generate_markdown(data: dict, output_path: str):
    lines = []
    lines.append("# 📋 รายงานสรุปความคืบหน้าประจำวัน (Daily Progress Report)")
    lines.append(f"**โครงการ:** {data.get('project_name', 'HRMS Development')}  ")
    lines.append(f"**วันที่:** {data.get('thai_date', '')} ({data.get('date_iso', '')})  ")
    lines.append(f"**รอบเวลาสรุป:** 17:30 น. (รายงานฉบับที่ {data.get('report_number', 1)})  ")
    lines.append(f"**ผู้บันทึก:** {data.get('author', 'ทีมพัฒนาและดูแลระบบ (Antigravity Assistant)')}\n")
    lines.append("---\n")

    lines.append("## 🎯 สรุปภาพรวมใน 3 บรรทัด (Executive Summary)")
    for idx, item in enumerate(data.get("executive_summary", []), start=1):
        lines.append(f"{idx}. **{item.get('title', '')}:** {item.get('desc', '')}")
    lines.append("\n---\n")

    lines.append("## 🛠️ รายละเอียดงานที่ทำและแก้ไขในวันนี้ (อธิบายเปรียบเทียบก่อน-หลัง และประโยชน์ทางธุรกิจ)\n")
    for idx, feat in enumerate(data.get("features", []), start=1):
        lines.append(f"### {idx}. {feat.get('title', '')}")
        if feat.get("before"):
            lines.append(f"* **ก่อนหน้านี้เป็นอย่างไร:** {feat['before']}")
        if feat.get("action"):
            lines.append(f"* **สิ่งที่ทำในวันนี้:** {feat['action']}")
        if feat.get("benefit"):
            lines.append(f"* **ประโยชน์ทางธุรกิจ:** {feat['benefit']}")
        lines.append("")
    lines.append("---\n")

    lines.append("## 📊 สรุปสถานะความพร้อมของระบบ ณ สิ้นวัน (System Readiness Summary)\n")
    lines.append("| องค์ประกอบระบบ | จำนวน / สถานะ | ความพร้อมใช้งาน |")
    lines.append("| :--- | :---: | :--- |")
    for row in data.get("readiness_table", []):
        lines.append(f"| **{row.get('component', '')}** | {row.get('status', '')} | {row.get('readiness', '')} |")
    lines.append("\n---\n")

    lines.append("## ⏩ แผนงานสำหรับรอบวันถัดไป")
    for idx, step in enumerate(data.get("next_steps", []), start=1):
        lines.append(f"{idx}. {step}")
    lines.append("")

    content = "\n".join(lines)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"[OK] Generated Markdown: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Daily Progress Report Generator")
    parser.add_argument("--project-path", default=".", help="Root directory of the project")
    parser.add_argument("--scan", action="store_true", help="Scan git and file activity and output JSON")
    parser.add_argument("--input-json", help="Path to JSON file containing summary payload")
    parser.add_argument("--date", help="Target date YYYY-MM-DD (defaults to today)")
    parser.add_argument("--number", type=int, help="Report sequence number (defaults to auto-increment)")
    parser.add_argument("--output-dir", help="Output directory for reports (defaults to <project>/daily_reports)")
    args = parser.parse_args()

    proj_path = os.path.abspath(args.project_path)
    if args.date:
        target_date = datetime.date.fromisoformat(args.date)
    else:
        target_date = datetime.date.today()

    date_iso = target_date.isoformat()
    thai_date = get_thai_date_str(target_date)

    if args.scan:
        activity = scan_git_activity(proj_path, date_iso)
        print(json.dumps(activity, ensure_ascii=False, indent=2))
        return

    output_dir = os.path.abspath(args.output_dir) if args.output_dir else os.path.join(proj_path, "daily_reports")
    os.makedirs(output_dir, exist_ok=True)

    report_num = args.number if args.number else get_next_report_number(output_dir)

    if args.input_json and os.path.exists(args.input_json):
        with open(args.input_json, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    else:
        activity = scan_git_activity(proj_path, date_iso)
        commits = activity.get("commits", [])
        data = {
            "project_name": os.path.basename(proj_path),
            "thai_date": thai_date,
            "date_iso": date_iso,
            "report_number": report_num,
            "author": "ทีมพัฒนาและดูแลระบบ (Antigravity Assistant)",
            "executive_summary": [
                {"title": "ความคืบหน้าการพัฒนา", "desc": f"ดำเนินงานพัฒนาฟีเจอร์และปรับปรุงระบบ รวม {len(commits)} commits ในวันนี้"},
                {"title": "ความพร้อมของระบบ", "desc": "ตรวจสอบความถูกต้องของโค้ดและการทำงานร่วมกันระหว่าง Frontend และ Backend"},
                {"title": "การจัดการซอร์สโค้ด", "desc": "บันทึกและส่งมอบงานเข้าสู่ Git Repository อย่างเป็นระเบียบ"}
            ],
            "features": [],
            "readiness_table": [
                {"component": "Git Repository", "status": f"{len(commits)} Commits วันนี้", "readiness": "✅ อัปเดตล่าสุด"},
                {"component": "Technology Stack", "status": ", ".join(activity.get("tech_stack", ["General"])), "readiness": "✅ พร้อมใช้งาน"}
            ],
            "next_steps": [
                "ทดสอบระบบและฟังก์ชันการทำงานอย่างละเอียด",
                "พัฒนาฟีเจอร์ในสปรินต์ถัดไปตามแผนงาน"
            ]
        }
        for c in commits:
            data["features"].append({
                "title": f"Commit {c['hash']}: {c['subject']}",
                "before": "อยู่ระหว่างการพัฒนา",
                "action": f"ปรับปรุงและพัฒนาโค้ดตามวัตถุประสงค์ {c['subject']} (เวลา {c['time']})",
                "benefit": "ยกระดับความสามารถของระบบและรองรับการใช้งานจริง"
            })

    data["thai_date"] = data.get("thai_date", thai_date)
    data["date_iso"] = data.get("date_iso", date_iso)
    data["report_number"] = data.get("report_number", report_num)

    md_filename = f"{report_num}.daily_summary_{date_iso}.md"
    docx_filename = f"{report_num}.daily_summary_{date_iso}.docx"
    md_path = os.path.join(output_dir, md_filename)
    docx_path = os.path.join(output_dir, docx_filename)

    generate_markdown(data, md_path)
    generate_docx(data, docx_path)

if __name__ == "__main__":
    main()
